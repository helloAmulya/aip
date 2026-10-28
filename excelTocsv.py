#! python3
# excelToCsv.py - Converts all Excel files in the current directory
# into separate CSV files for each worksheet.

import os
import csv
import openpyxl

for excelFile in os.listdir("."):
    # Skip files that are not Excel files.
    if not excelFile.endswith(".xlsx"):
        continue

    # Load the workbook.
    wb = openpyxl.load_workbook(excelFile, data_only=True)

    # Loop through every worksheet.
    for sheetName in wb.sheetnames:
        sheet = wb[sheetName]

        # Create the CSV filename.
        excelFilename = os.path.splitext(excelFile)[0]
        csvFilename = f"{excelFilename}_{sheetName}.csv"

        # Create the CSV file and writer object.
        with open(csvFilename, "w", newline="") as csvFile:
            csvWriter = csv.writer(csvFile)

            # Loop through every row in the worksheet.
            for row in sheet.iter_rows(values_only=True):
                rowData = list(row)

                # Write the row to the CSV file.
                csvWriter.writerow(rowData)

        print(f"Converted: {excelFile} [{sheetName}] -> {csvFilename}")

    wb.close()

print("Conversion complete!")
