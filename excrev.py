#! python3

from openpyxl import Workbook

school_wb = Workbook()  # initiate the workbook

# if you need to work on an individual sheet in a book, use 'school_wb[worksheet_name]'

ws1 = school_wb.active  # set sheet 1 as active (default)

# school_wb.remove(school_wb.active) # Remove the default sheet

ws1.title = "Students"  # add title for sheet1

# create other sheets
ws2 = school_wb.create_sheet("Teachers")
ws3 = school_wb.create_sheet("Subjects")

# add data to each sheet

ws1.append(["Name", "Age"])  # default header row
ws1.append(["alice", "14"])  # data gets added after this

ws2.append(["Name", "Department"])
ws2.append(["mathew", "science"])

ws3.append(["Subject", "Code"])
ws3.append(["python", "404"])


# save the workbok
school_wb.save("School.xlsx")

# imp -> using 'school_wb = Workbook()' instead of openpyxl.load_workbook(), and saving it will rewrite the exisiting file, for adding new data use load_workbook()
