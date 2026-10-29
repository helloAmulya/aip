#! python3
# updateProduce.py - Corrects costs in produce sales spreadsheet.


import openpyxl

wb = openpyxl.load_workbook("produceSales.xlsx")

sheet = wb["Sheet"]


# products and their updated prices
price_update = {"Garlic": 3.07, "Celery": 1.19, "Lemon": 1.27}

count = dict.fromkeys(
    price_update, 0
)  # take the products from price_update, avoid repeating product names

for rowNum in range(2, sheet.max_row + 1):
    oldPrice = sheet.cell(row=rowNum, column=2).value
    produceName = sheet.cell(row=rowNum, column=1).value
    if produceName in price_update:
        count[produceName] += 1
        sheet.cell(row=rowNum, column=2).value = price_update[produceName]
        # print(
        #     "updating %s with %s price to %s"
        #     % (produceName,oldPrice, price_update[produceName])
        # )
print("values updated: ", count)
wb.save("updatedProduceSales.xlsx")
